"""Report generation: marked-up photo, zoomed snippets (expected vs found), filled Excel checklist,
GA overlay for the rectified view, and re-check (before vs after) with issue history.
The PDF report lives in ``pdf_report.py``."""
import datetime
import json
import os
import re

import cv2
import numpy as np
import pymupdf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont

import config
from inspector import R

F = config.font("DejaVuSans-Bold.ttf")
FR = config.font("DejaVuSans.ttf")
COL = {"FAIL": (225, 29, 72), "REVIEW": (245, 158, 11), "PASS": (16, 185, 129), "SKIP": (120, 120, 130),
       "NOT_IN_VIEW": (100, 116, 139), "CLOSED": (16, 185, 129), "OPEN": (225, 29, 72), "REGRESSION": (147, 51, 234),
       "NEW": (234, 88, 12)}
ISSUE = ("FAIL", "REVIEW", "CLOSED", "OPEN", "REGRESSION", "NEW")
DRAWING_BLUE = (31, 79, 168)


def to_photo(box, H):
    """Drawing box (mm) -> 4 corner points in photo pixels."""
    x, y, w, h = box
    pts = np.float32([[x, y], [x + w, y], [x + w, y + h], [x, y + h]]) * R
    return cv2.perspectiveTransform(pts[None], np.linalg.inv(np.array(H)))[0]


def numbered(results):
    k = 0
    for r in results:
        if r["result"] in ("FAIL", "REVIEW"):
            k += 1; r["marker"] = k
    return results


def status_of(r, status_key="result"):
    return r.get(status_key) or r["result"]


def markup(photo_path, rep, out_path, status_key="result", scale=1.0):
    """Box and number every issue on the photo (translucent fill, thick outline, numbered badge).

    ``scale`` maps the report's photo pixels onto the image given (a preview smaller than the inspected photo).
    """
    im = Image.open(photo_path).convert("RGB")
    over = Image.new("RGBA", im.size, (0, 0, 0, 0)); d = ImageDraw.Draw(over)
    u = max(0.35, min(1.0, max(im.size) / 2900))         # badge and caption sizes follow the image size
    f = ImageFont.truetype(F, max(14, int(46 * u))); fs = ImageFont.truetype(F, max(11, int(26 * u)))
    for r in rep["results"]:
        st = status_of(r, status_key)
        if st not in ISSUE or "marker" not in r or r.get("marker") is None:
            continue
        p = to_photo(r["box"], rep["H"]) * scale; c = COL[st]
        pad = 10 * u
        x0, y0 = p[:, 0].min() - pad, p[:, 1].min() - pad
        x1, y1 = p[:, 0].max() + pad, p[:, 1].max() + pad
        d.rectangle([x0, y0, x1, y1], fill=c + (46,), outline=c + (255,), width=max(3, int(9 * u)))
        cap = f"{st} · {r['id']}"
        tw = d.textlength(cap, font=fs)
        ty = y1 + 8 * u if y1 + 50 * u < im.height else y0 - 46 * u
        d.rounded_rectangle([x0, ty, x0 + tw + 24 * u, ty + 38 * u], radius=10 * u, fill=(10, 14, 20, 225))
        d.text((x0 + 12 * u, ty + 5 * u), cap, fill=c + (255,), font=fs)
        cx, cy, rr = x1 + 4 * u, y0 - 34 * u, 36 * u
        d.ellipse([cx - rr, cy - rr, cx + rr, cy + rr], fill=c + (255,), outline=(255, 255, 255, 255),
                  width=max(2, int(5 * u)))
        t = str(r["marker"]); w_ = d.textlength(t, font=f)
        d.text((cx - w_ / 2, cy - 28 * u), t, fill=(255, 255, 255, 255), font=f)
    im = Image.alpha_composite(im.convert("RGBA"), over).convert("RGB")
    im.save(out_path, quality=90)
    return out_path


# ------------------------------------------------------------------ drawing rasters and crops
def ga_png(ga_pdf, out_png, dpi=150):
    """Rasterise page 1 of the GA; returns the mm->px mapping of the backplate (ox, oy, k px/mm)."""
    with pymupdf.open(ga_pdf) as doc:
        page = doc[0]
        page.get_pixmap(dpi=dpi).save(out_png)
        rects = [pymupdf.Rect(it[1]) for dr in page.get_drawings() for it in dr["items"] if it[0] == "re"]
        m = re.search(r"BACKPLATE\s+(\d+)\s*x\s*(\d+)\s*mm", page.get_text())
    bp = max(rects, key=lambda r: r.width * r.height)
    z = dpi / 72
    width_mm = int(m.group(1)) if m else 600
    return dict(ox=bp.x0 * z, oy=bp.y0 * z, k=bp.width * z / width_mm)


def ga_overlay(ga_pdf, size_mm, out_png, px_per_mm=2, rgb=(56, 189, 248)):
    """Transparent PNG of the drawing linework, cropped to the backplate, aligned to the rectified photo."""
    tmp = out_png + ".tmp.png"
    s = ga_png(ga_pdf, tmp, dpi=220)
    g = Image.open(tmp).convert("L"); os.remove(tmp)
    W, H = size_mm
    box = (int(round(s["ox"])), int(round(s["oy"])), int(round(s["ox"] + W * s["k"])), int(round(s["oy"] + H * s["k"])))
    g = g.crop(box).resize((W * px_per_mm, H * px_per_mm), Image.LANCZOS)
    a = 255 - np.asarray(g, np.float32)
    a = np.clip(a * 1.6, 0, 255).astype(np.uint8)
    out = np.zeros((a.shape[0], a.shape[1], 4), np.uint8); out[..., :3] = rgb; out[..., 3] = a
    Image.fromarray(out, "RGBA").save(out_png, optimize=True)
    return out_png


def _outline(img, pts, color, width):
    d = ImageDraw.Draw(img)
    d.polygon([tuple(map(float, p)) for p in pts], outline=color, width=width)
    return img


def expected_crop(ga_img, s, box, margin=15, color=(14, 165, 233)):
    """Crop of the GA drawing around an item (the drawing is the ground truth)."""
    x, y, w, h = box
    X0, Y0 = s["ox"] + (x - margin) * s["k"], s["oy"] + (y - margin) * s["k"]
    X1, Y1 = s["ox"] + (x + w + margin) * s["k"], s["oy"] + (y + h + margin) * s["k"]
    c = ga_img.crop((int(X0), int(Y0), int(X1), int(Y1))).convert("RGB")
    k = s["k"]
    pts = [(margin * k, margin * k), ((margin + w) * k, margin * k), ((margin + w) * k, (margin + h) * k), (margin * k, (margin + h) * k)]
    return _outline(c, pts, color, max(2, int(k * 0.6)))


def found_crop(photo_img, H, box, color, margin=60, scale=1.0, max_side=None):
    """Zoomed crop of the photo around where the drawing says the item is (``scale``: image px per report px)."""
    p = to_photo(box, H) * scale
    margin = max(12, margin * scale)
    x0, y0 = int(p[:, 0].min() - margin), int(p[:, 1].min() - margin)
    x1, y1 = int(p[:, 0].max() + margin), int(p[:, 1].max() + margin)
    c = photo_img.crop((x0, y0, x1, y1)).convert("RGB")
    c = _outline(c, p - np.float32([x0, y0]), color, max(2, int(5 * min(1.0, scale * 2))))
    if max_side and max(c.size) > max_side:
        c.thumbnail((max_side, max_side), Image.LANCZOS)
    return c


def aligned_crop(rect_img, box, color, margin=15, k=R):
    """Same window as ``expected_crop`` but from the perspective-corrected photo (pixel-aligned with the drawing);
    ``k`` = px per mm of ``rect_img`` (stored rectified images are kept at a reduced resolution)."""
    x, y, w, h = box
    c = rect_img.crop((int((x - margin) * k), int((y - margin) * k), int((x + w + margin) * k), int((y + h + margin) * k))).convert("RGB")
    pts = [(margin * k, margin * k), ((margin + w) * k, margin * k), ((margin + w) * k, (margin + h) * k), (margin * k, (margin + h) * k)]
    return _outline(c, pts, color, max(2, int(k)))


def snippet(photo_path, rep, r, ga_png_path, ga_scale, out_path):
    """Expected (crop of the GA drawing) next to Found (zoomed photo crop), on one canvas."""
    st = status_of(r, "recheck") if r.get("recheck") else r["result"]
    found = found_crop(Image.open(photo_path), rep["H"], r["box"], COL.get(st, COL["FAIL"]))
    exp = expected_crop(Image.open(ga_png_path), ga_scale, r["box"])
    Ht = 300
    found = found.resize((max(1, int(found.width * Ht / found.height)), Ht))
    exp = exp.resize((max(1, int(exp.width * Ht / exp.height)), Ht))
    W_ = found.width + exp.width + 30
    canvas = Image.new("RGB", (W_, Ht + 64), (255, 255, 255)); dd = ImageDraw.Draw(canvas)
    fnt = ImageFont.truetype(F, 22)
    canvas.paste(exp, (0, 50)); canvas.paste(found, (exp.width + 30, 50))
    dd.text((4, 14), "EXPECTED (GA)", fill=DRAWING_BLUE, font=fnt)
    dd.text((exp.width + 34, 14), "FOUND (photo)", fill=COL.get(st, COL["FAIL"]), font=fnt)
    canvas.save(out_path)
    return out_path


# ------------------------------------------------------------------ Excel checklist
def fill_checklist(template_xlsx, rep, out_xlsx, meta=None):
    """Fill Emerson's QC checklist template in place (Expected / Found / Result / Evidence) and add a summary sheet."""
    wb = load_workbook(template_xlsx); ws = wb.active
    by = {r["id"]: r for r in rep["results"]}
    fills = {"PASS": "DCFCE7", "FAIL": "FECACA", "REVIEW": "FEF3C7", "SKIP": "E5E7EB", "NOT_IN_VIEW": "F1F5F9",
             "CLOSED": "DCFCE7", "OPEN": "FECACA", "REGRESSION": "E9D5FF", "NEW": "FFEDD5"}
    is_recheck = any(r.get("recheck") for r in rep["results"]) or bool(meta and meta.get("parent"))
    thin = Side(style="thin", color="D1D5DB")
    if is_recheck:
        ws.cell(row=1, column=8, value="Re-check")
    for row in ws.iter_rows(min_row=2):
        r = by.get(row[0].value)
        if not r:
            continue
        row[3].value = r["expected"]; row[4].value = r["found"]; row[5].value = r["result"]
        row[6].value = (f"Marker #{r['marker']} | " if r.get("marker") else "") + "; ".join(map(str, r["evidence"]))[:200]
        row[5].fill = PatternFill("solid", fgColor=fills[r["result"]]); row[5].font = Font(bold=True)
        row[5].alignment = Alignment(horizontal="center")
        if is_recheck and r.get("recheck"):
            c = ws.cell(row=row[0].row, column=8, value=r["recheck"])
            c.fill = PatternFill("solid", fgColor=fills[r["recheck"]]); c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for c in row:
            c.border = Border(bottom=thin)
    for c, wdt in zip("ABCDEFGH", [16, 10, 34, 22, 34, 10, 60, 13]):
        ws.column_dimensions[c].width = wdt
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F2742")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary", 0)
    k = {x: sum(r["result"] == x for r in rep["results"]) for x in ("PASS", "FAIL", "REVIEW", "SKIP", "NOT_IN_VIEW")}
    verdict = (meta or {}).get("verdict", {}).get("label") or ("HOLD" if k["FAIL"] else "REVIEW" if k["REVIEW"] else "PASS")
    rows = [("EAGLE EYE · Cabinet QC", ""), ("", ""),
            ("Cabinet", rep.get("ga", {}).get("cabinet", "")),
            ("Photo", os.path.basename((meta or {}).get("photo_name") or rep["photo"])),
            ("Inspected", (meta or {}).get("created", datetime.datetime.now().isoformat(timespec="seconds"))),
            ("Verdict", verdict), ("", ""),
            ("PASS", k["PASS"]), ("FAIL", k["FAIL"]), ("REVIEW", k["REVIEW"]), ("SKIP (root cause elsewhere)", k["SKIP"])]
    if k["NOT_IN_VIEW"] or rep.get("engine_kind") == "real":
        cov = rep.get("coverage") or {}
        rows += [("NOT_IN_VIEW (not in this photo)", k["NOT_IN_VIEW"]),
                 ("Coverage", f"{cov.get('verified', 0)} of {cov.get('total', 0)} labelled drawing items verified")]
    rows += [("Items checked", len(rep["results"]))]
    rc = (meta or {}).get("recheck")
    if rc:
        rows += [("", ""), ("Re-check of run", (meta or {}).get("parent", ""))] + [(x, rc.get(x, 0)) for x in ("CLOSED", "OPEN", "REGRESSION", "NEW")]
    for r_ in rows:
        s.append(list(r_))
    s["A1"].font = Font(bold=True, size=16, color="0F2742")
    for row in s.iter_rows(min_row=3):
        row[0].font = Font(bold=True, color="374151")
        if row[0].value in fills:
            row[1].fill = PatternFill("solid", fgColor=fills[row[0].value]); row[1].font = Font(bold=True)
    vc = s["B6"]; vc.font = Font(bold=True, size=13, color="FFFFFF")
    vc.fill = PatternFill("solid", fgColor={"HOLD": "E11D48", "REVIEW": "D97706"}.get(verdict, "059669"))
    s.column_dimensions["A"].width = 30; s.column_dimensions["B"].width = 40
    wb.active = 1
    wb.save(out_xlsx)
    return out_xlsx


CHECK_NAMES = {"GA": "Part present", "MISSING": "Part present", "ORDER": "Order", "LABEL": "Device label",
               "UNKNOWN_LABEL": "Unknown label", "FERRULE": "Wire ferrule", "COLOUR": "Wire colour",
               "REGISTRATION": "Drawing alignment", "CHECKLIST": "Manual checklist requirements"}


def checklist_template(rep, out_xlsx):
    """QC checklist generated from the drawing's knowledge base when the user supplies none (real drawing sets):
    one row per checklist item of the report, in the column layout ``fill_checklist`` fills."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "QC Checklist"
    ws.append(["Item", "Check", "Description", "Expected (drawing)", "Found (photo)", "Result", "Evidence"])
    for r in rep["results"]:
        view = f" ({r['view']} view)" if r.get("view") else ""
        desc = {"MISSING": f"{r['target']} installed where the GA drawing places it{view}",
                "LABEL": f"Label of {r['target']} reads exactly '{r['target']}'{view}",
                "ORDER": f"Devices of group {r['target']} in the drawing's order",
                "UNKNOWN_LABEL": f"Label '{r['target']}' is a tag of the drawing set",
                "FERRULE": f"Ferrule {r['target']} names a connection of the drawing"}.get(r["check"], r["check"])
        ws.append([r["id"], CHECK_NAMES.get(r["check"], r["check"]), desc, "", "", "", ""])
    wb.save(out_xlsx)
    return out_xlsx


# ------------------------------------------------------------------ re-check and history
def recheck(prev, new):
    """Compare a new inspection with the previous one.

    CLOSED     was FAIL/REVIEW, now PASS (fixed)
    OPEN       was FAIL/REVIEW, still not PASS
    REGRESSION was PASS, now FAIL (the fix broke something)
    NEW        newly flagged (e.g. a label that was SKIPped while its part was missing)
    Items keep the marker number they had in the previous report; new issues get the next numbers.
    """
    pb = {r["id"]: r for r in prev["results"]}
    nxt = max([r.get("marker") or 0 for r in prev["results"]] + [0])
    for r in new["results"]:
        r.pop("marker", None)
        p = pb.get(r["id"])
        if p and p["result"] in ("FAIL", "REVIEW"):
            r["recheck"] = "CLOSED" if r["result"] == "PASS" else "OPEN"
            if p.get("marker"):
                r["marker"] = p["marker"]
            else:
                nxt += 1; r["marker"] = nxt
        elif r["result"] in ("FAIL", "REVIEW"):
            r["recheck"] = "REGRESSION" if (p and p["result"] == "PASS" and r["result"] == "FAIL") else "NEW"
            nxt += 1; r["marker"] = nxt
        else:
            continue
        if p:
            r["prev_result"], r["prev_found"] = p["result"], p["found"]
    new["recheck"] = {k: sum(r.get("recheck") == k for r in new["results"]) for k in ("CLOSED", "OPEN", "REGRESSION", "NEW")}
    return new


def history_append(path, rep, label):
    h = json.load(open(path)) if os.path.exists(path) else []
    h.append(dict(time=datetime.datetime.now().isoformat(timespec="seconds"), run=label, photo=os.path.basename(rep["photo"]),
                  fail=sum(r["result"] == "FAIL" for r in rep["results"]),
                  passed=sum(r["result"] == "PASS" for r in rep["results"]),
                  issues=[r["id"] for r in rep["results"] if r["result"] in ("FAIL", "REVIEW")]))
    json.dump(h, open(path, "w"), indent=1)
