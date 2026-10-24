"""Inspection service shared by the web console (server.py) and the Streamlit app.

A *run* is one inspection of one photo, stored in ``out/runs/<run_id>/`` as metadata, never as a copy of the photo:

    report.json           every checklist item: expected, found, result, evidence, box (photo px or drawing mm)
    meta.json             summary used by history lists, incl. the photo's metadata (SHA-256 fingerprint, size,
                          EXIF orientation, quality) - the photo itself is read once and discarded
    preview.jpg           small display copy (<= 1600 px) the console draws the boxes on
    marked.jpg            the preview with every issue boxed and numbered
    rectified.jpg         photo warped onto the drawing (demo: 2 px/mm; real cabinet: straightened onto the primary
                          GA view through the fitted perspective)
    ga_overlay.png        GA linework aligned with rectified.jpg
    drawing_overlay.png   real cabinet: the fitted GA views' linework laid onto the photo through the perspective
    crops/                expected-vs-found proof images (issue crops are cut from the full-resolution photo while
                          it is in memory; the others are made on demand from the preview / the drawing)
    QC_Checklist_FILLED.xlsx, EAGLE_EYE_QC_Report.pdf

Drawing sets and checklists are kept once, content-addressed, under ``cache/inputs/<sha256>``. Engine results are
cached by fingerprint (photo SHA-256 + input SHA-256s + checks + engine version + verifier), so inspecting the same photo
again needs neither the photo nor the engine: ``cached_result`` answers from the metadata alone.

A re-check run points at its parent; the chain (v1 -> v2 -> v3) is the issue history.
"""
import datetime
import hashlib
import json
import os
import random
import re
import shutil
import uuid
from pathlib import Path

import numpy as np
from PIL import Image, ImageOps

import assembly
import calibration
import config
import pdf_report
import report as rp
from inspector import CHECKS, ORDER, R, InspectionError, counts, inspect, unregistered_results
from render import describe_fault, expected_item, random_faults, render

FILES = dict(marked="marked.jpg", rectified="rectified.jpg", overlay="ga_overlay.png", json="report.json",
             checklist="QC_Checklist_FILLED.xlsx", pdf="EAGLE_EYE_QC_Report.pdf", preview="preview.jpg",
             drawing="drawing_overlay.png")
CACHE = config.DATA / "cache"
CACHE.mkdir(parents=True, exist_ok=True)
INPUTS = CACHE / "inputs"                  # drawing sets / wire lists / checklists, content-addressed (sha256)
RESULTS = CACHE / "results"                # engine results by fingerprint (no photos)
WORK = CACHE / "work"                      # a photo lives here only while it is being inspected
PREVIEW_SIDE = 1600
RECT_PX_PER_MM = 2                         # stored rectified image resolution (the engine works at inspector.R)
CROP_SIDE = 520                            # stored proof crops are at most this many px per side
RESULT_SCHEMA = "eagle-eye-result/2"  # invalidates results made before registration/approval guards
SAFETY_SCHEMA = "eagle-eye-evidence/1"
_ID = re.compile(r"^[0-9]{8}-[0-9]{6}-[0-9a-f]{4}$")
_SHA = re.compile(r"^[0-9a-f]{64}$")


# ------------------------------------------------------------------ helpers
def now():
    return datetime.datetime.now().isoformat(timespec="seconds")


def new_id():
    return datetime.datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:4]


def run_dir(rid):
    if not _ID.match(str(rid)):
        raise KeyError(rid)
    d = config.RUNS / rid
    if not d.exists():
        raise KeyError(rid)
    return d


def _sha(path):
    return hashlib.sha1(Path(path).read_bytes()).hexdigest()[:16]


def _safe(item_id):
    return re.sub(r"[^A-Za-z0-9]+", "_", item_id)


def write_json(path, obj):
    tmp = Path(str(path) + ".tmp")
    tmp.write_text(json.dumps(obj, indent=1, default=str, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def photo_record(src, name=None):
    """Metadata of a photo (what is kept instead of the photo): SHA-256 of the file bytes, upright size, EXIF
    orientation, byte size and file name. Raises InspectionError for unreadable images."""
    src = Path(src)
    try:
        with Image.open(src) as im:
            im.load()
            orient = int(im.getexif().get(0x0112, 1) or 1)
            w, h = im.size
            fmt = im.format
    except Exception as e:  # noqa: BLE001
        raise InspectionError("Could not read the cabinet photo. Upload a JPG or PNG image.") from e
    if orient in (5, 6, 7, 8):
        w, h = h, w
    return dict(sha256=sha256_file(src), bytes=src.stat().st_size, name=name or src.name, size=[w, h],
                exif_orientation=orient, format=fmt)


def prepare_photo(src, dst_dir):
    """Working copy of the photo for the demo engine: EXIF rotation applied, huge phone photos brought to the tuned
    size. It lives in a temporary work folder and is deleted after the inspection."""
    src = Path(src)
    try:
        im = Image.open(src)
        im.load()
    except Exception as e:  # noqa: BLE001
        raise InspectionError("Could not read the cabinet photo. Upload a JPG or PNG image.") from e
    orient = im.getexif().get(0x0112, 1)
    big = max(im.size) > 3400
    if orient == 1 and not big and im.format in ("JPEG", "PNG"):
        dst = Path(dst_dir) / ("photo.png" if im.format == "PNG" else "photo.jpg")
        shutil.copyfile(src, dst)
        return dst
    im = ImageOps.exif_transpose(im).convert("RGB")
    if big:
        im.thumbnail((2900, 2900), Image.LANCZOS)
    dst = Path(dst_dir) / "photo.jpg"
    im.save(dst, quality=95)
    return dst


def save_preview(rgb, dst, side=PREVIEW_SIDE):
    """Small display copy of an upright RGB array; returns (path, scale = preview px / photo px)."""
    im = Image.fromarray(np.ascontiguousarray(rgb))
    k = min(1.0, side / max(im.size))
    if k < 1:
        im = im.resize((max(1, round(im.width * k)), max(1, round(im.height * k))), Image.LANCZOS)
    im.save(dst, quality=84, optimize=True, progressive=True)
    return Path(dst), k


def ga_raster(ga_pdf, dpi=200):
    """Raster of the GA drawing for proof crops, cached by file content."""
    key = _sha(ga_pdf)
    png, js = CACHE / f"ga_{key}_{dpi}.png", CACHE / f"ga_{key}_{dpi}.json"
    if not (png.exists() and js.exists()):
        s = rp.ga_png(str(ga_pdf), str(png), dpi=dpi)
        write_json(js, s)
    return png, read_json(js)


def thumbnail(path, width):
    """Cached JPEG thumbnail of any image (history cards, sample gallery)."""
    path = Path(path)
    width = int(max(64, min(1600, width)))
    key = hashlib.sha1(f"{path.resolve()}|{path.stat().st_mtime_ns}|{width}".encode()).hexdigest()[:20]
    out = CACHE / "thumbs" / f"{key}.jpg"
    if not out.exists():
        out.parent.mkdir(exist_ok=True)
        im = Image.open(path).convert("RGB")
        im.thumbnail((width, width * 4), Image.LANCZOS)
        im.save(out, quality=84)
    return out


class NeedPhoto(InspectionError):
    """The engine has no stored result for this photo fingerprint: the photo itself has to be sent."""


class WrongDrawing(InspectionError):
    """Not one label of the drawing set was found in the photo: it is another cabinet's photo (or unreadable).
    Asking for the right drawing set beats a report in which every item is REVIEW."""


DEMO_MISMATCH = ("None of the {n} labels of the demo drawing {cab} could be matched in this photo, so its cabinet identity "
                 "and alignment could not be verified. "
                 "Add this cabinet's own drawing set (vector PDF) under Engineering inputs - or pick the real-cabinet "
                 "preset if one is offered - and inspect again.")
REAL_MISMATCH = ("None of the labels of drawing set {name} ({cab}) could be read in this photo. Check that the photo "
                 "shows this cabinet, then re-photo one side at a time from about 1 m, straight on, so the device labels "
                 "are sharp.")
DRAWING_REQUIRED = ("Add this cabinet's matching GA drawing PDF under Engineering inputs, or select its real-cabinet "
                    "preset. Uploaded photographs are never inspected against the synthetic demo drawing automatically. "
                    "For a generated test image, explicitly select Synthetic demo mode.")


def validate_input(path, kind):
    """Reject unreadable engineering inputs before OCR or background work starts."""
    try:
        if kind == "ga":
            import pymupdf
            with pymupdf.open(path) as doc:
                if not doc.is_pdf or doc.needs_pass or not doc.page_count:
                    raise ValueError("empty, encrypted or non-PDF document")
                if not any(page.get_text().strip() for page in doc):
                    raise InspectionError("The drawing PDF has no selectable text. Upload the original vector GA PDF; "
                                          "a scanned drawing or photograph cannot provide reliable cabinet requirements.")
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            wb.close()
    except InspectionError:
        raise
    except Exception as exc:  # noqa: BLE001 - translate parser errors to an actionable upload error
        title = "GA drawing PDF" if kind == "ga" else "wire list workbook" if kind == "wires" else "checklist workbook"
        raise InspectionError(f"Could not read the {title}. Upload a valid, unencrypted {'PDF' if kind == 'ga' else 'XLSX'} file.") from exc


def validate_checks(enabled, kind=None):
    """None selects defaults; an explicitly empty or misspelled selection is an input error."""
    if enabled is None:
        return
    from checks import ORDER as REAL_ORDER
    allowed = set(ORDER)
    if kind != "legacy":
        allowed.update(REAL_ORDER)
    if not isinstance(enabled, (list, tuple)) or not enabled:
        raise InspectionError("Select at least one inspection check, or omit checks to use all required defaults.")
    unknown = sorted({str(c) for c in enabled if not isinstance(c, str) or c not in allowed})
    if unknown:
        raise InspectionError(f"Unknown inspection check(s): {', '.join(unknown)}. Supported checks: {', '.join(sorted(allowed))}.")


def validate_request(ga=None, engine="auto", source="upload", prev=None, photo=None, enabled=None):
    """An arbitrary photograph must never silently opt into the demonstration drawing."""
    if engine not in ("auto", "legacy", "real"):
        raise InspectionError("engine must be auto, legacy (synthetic demo), or real.")
    validate_checks(enabled)
    sample_photo = bool(photo and Path(photo).resolve().parent == config.SAMPLES.resolve())
    demo_session = bool(prev and prev.get("source") in ("sample", "forge") and prev.get("engine_kind", "legacy") == "legacy")
    if not ga and source not in ("sample", "forge") and engine != "legacy" and not sample_photo and not demo_session:
        raise InspectionError(DRAWING_REQUIRED)


def verdict(c, n_items, warnings=None, rep=None):
    real = bool(rep and rep.get("engine_kind") == "real")
    weak = next((w for w in warnings or [] if w["code"] == "weak_registration"), None)
    if weak:
        if real:
            return dict(label="REVIEW", headline="REVIEW · photo not placed on the drawing",
                        detail=weak["message"] + " Re-photo one side of the cabinet at a time, straight on.")
        return dict(label="REVIEW", headline="REVIEW · photo not aligned to the drawing",
                    detail="The photo could not be aligned with the drawing. Component checks were not verified. "
                           "Confirm the matching cabinet drawing, then photograph the required area straight on.")
    if c["FAIL"]:
        return dict(label="HOLD", headline=f"HOLD · {c['FAIL']} fault{'s' if c['FAIL'] != 1 else ''} to fix before shipping",
                    detail="Every FAIL is boxed on the photo with expected-vs-found proof."
                           + (f" {c['REVIEW']} check(s) also require review." if c["REVIEW"] else ""))
    if c["REVIEW"]:
        return dict(label="REVIEW", headline=f"REVIEW · {c['REVIEW']} unresolved check{'s' if c['REVIEW'] != 1 else ''}",
                    detail="No confirmed faults. Review the evidence and each check's guidance: a matching drawing, "
                           "additional photographs, or a manual check may be required before approval.")
    if not c.get("PASS") or c.get("SKIP") or c.get("NOT_IN_VIEW"):
        return dict(label="REVIEW", headline="REVIEW · inspection coverage incomplete",
                    detail="Cabinet approval is blocked: required checks were skipped, outside the photograph, or no checks were verified.")
    if rep and set(rep.get("enabled") or []) != set(_enabled_for("real" if real else "legacy", None)):
        return dict(label="REVIEW", headline="REVIEW · required checks disabled",
                    detail="The selected checks passed, but all required inspection checks must be enabled before approval.")
    if real:
        cov = rep.get("coverage") or {}
        total = cov.get("total")
        if not (cov.get("approval_allowed") is True and cov.get("complete") is True and
                isinstance(total, (int, float)) and total > 0 and (cov.get("verified") or 0) >= total):
            return dict(label="REVIEW", headline="REVIEW · inspection coverage incomplete",
                        detail=f"This photograph verified {cov.get('verified', 0)} of {cov.get('total', 0)} drawing items. "
                               "Capture the remaining required checks before cabinet approval.")
        return dict(label="PASS", headline="PASS · required visual checks verified",
                    detail=f"All {n_items} required visual checks passed against the supplied drawing.")
    return dict(label="PASS", headline="PASS · selected visual checks verified",
                detail=f"All {n_items - c['SKIP']} items the drawing expects were verified on the photo.")


# ------------------------------------------------------------------ inputs
def default_inputs():
    return dict(ga=config.DEMO_GA, wires=config.DEMO_WIRES, checklist=config.DEMO_CHECKLIST)


def store_input(src, name, sha=None):
    """Keep an input file once, content-addressed (cache/inputs/<sha256><ext>); returns its meta reference."""
    sha = sha or sha256_file(src)
    ext = Path(name).suffix.lower() or Path(src).suffix.lower()
    dst = INPUTS / f"{sha}{ext}"
    if not dst.exists():
        INPUTS.mkdir(parents=True, exist_ok=True)
        tmp = dst.with_suffix(dst.suffix + f".{uuid.uuid4().hex[:6]}.tmp")
        shutil.copyfile(src, tmp)
        os.replace(tmp, dst)
    return dict(path=str(dst), name=name, demo=False, sha256=sha, ext=ext)


def input_by_sha(sha, ext):
    """Path of a stored input with this SHA-256 (fingerprint-only requests), or None."""
    if not _SHA.match(str(sha or "")) or ext not in (".pdf", ".xlsx"):
        return None
    p = INPUTS / f"{sha}{ext}"
    return p if p.exists() else None


def input_path(rid, meta, key):
    """Where a run's input file lives now (robust to the project folder being moved or copied); None if not given."""
    ref = meta["inputs"][key]
    if ref.get("missing"):
        return None
    if ref["demo"]:
        return default_inputs()[key]
    if ref.get("sha256"):
        p = INPUTS / f"{ref['sha256']}{ref.get('ext') or Path(ref['name']).suffix.lower()}"
        if p.exists():
            return p
    if ref.get("path"):
        try:
            return run_dir(rid) / "inputs" / Path(ref["path"]).name    # runs stored before content addressing
        except KeyError:
            return None
    return None


def resolve_inputs(d, ga=None, wires=None, checklist=None, names=None, defaults=True):
    """Uploaded files when given (stored content-addressed), otherwise the demo inputs (demo engine) or nothing."""
    names = names or {}
    out = {}
    for key, given, default in (("ga", ga, config.DEMO_GA), ("wires", wires, config.DEMO_WIRES),
                                ("checklist", checklist, config.DEMO_CHECKLIST)):
        if given:
            out[key] = store_input(given, names.get(key) or Path(given).name)
        elif defaults:
            out[key] = dict(path=str(default), name=default.name, demo=True)
        else:
            out[key] = dict(path=None, name="—", demo=False, missing=True)
    return out


def inherited_inputs(prev, ga=None, wires=None, checklist=None, names=None):
    """Reuse unchanged engineering inputs; a replacement GA starts a new reference set."""
    names = dict(names or {})
    if not prev or ga:
        return ga, wires, checklist, names
    supplied = {"ga": ga, "wires": wires, "checklist": checklist}
    for key, ref in (prev.get("inputs") or {}).items():
        if key not in supplied or supplied[key] or ref.get("demo") or ref.get("missing"):
            continue
        supplied[key] = input_path(prev.get("id"), prev, key)
        if not supplied[key] or not Path(supplied[key]).is_file():
            raise InspectionError(f"The original {key} input is no longer available. Upload it again for this re-check.")
        names.setdefault(key, ref.get("name") or Path(supplied[key]).name)
    return supplied["ga"], supplied["wires"], supplied["checklist"], names


def engine_for(ga, engine="auto", prev=None):
    """A newly supplied drawing takes priority over the previous run's engine."""
    if engine in ("legacy", "real"):
        return engine
    if ga:
        import pipeline
        return "real" if pipeline.is_real_drawing(ga) else "legacy"
    if prev:
        return prev
    return "legacy"


def _enabled_for(kind, enabled):
    validate_checks(enabled, kind)
    if kind == "legacy":
        return list(ORDER) if enabled is None else list(dict.fromkeys(enabled))
    import pipeline
    return pipeline.real_enabled(enabled)


# ------------------------------------------------------------------ result cache (fingerprints, no photos)
def _agent_signature(kind):
    if kind != "real":
        return None
    try:
        from agent import status
        from consensus import fail_without_agent
        st = status()
        return [st.get("enabled"), st.get("backend"), st.get("model"), st.get("fast_model"), fail_without_agent()]
    except Exception:  # noqa: BLE001
        return ["unavailable"]


def result_key(photo_sha, inputs, enabled, kind):
    """Fingerprint of one inspection: same photo + same inputs + same checks + same engine = same result."""
    payload = dict(schema=RESULT_SCHEMA, photo=photo_sha, kind=kind, version=config.VERSION, enabled=sorted(enabled),
                   inputs={k: (v.get("sha256") or (f"demo:{v['name']}" if v.get("demo") else None))
                           for k, v in inputs.items()},
                   agent=_agent_signature(kind), ocr=os.environ.get("EAGLE_OCR_MODE", ""),
                   colour=config.COLOUR_PROFILE or "")
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def load_result(key):
    d = RESULTS / key if key else None
    if not d or not (d / "result.json").exists():
        return None
    try:
        data = read_json(d / "result.json")
    except (OSError, ValueError):
        return None
    if data.get("schema") != RESULT_SCHEMA or not (d / FILES["preview"]).exists():
        return None
    data["dir"] = d
    return data


def store_result(key, rep, art, photo_meta, preview_scale, events):
    """Cache an engine result (report + preview + proof crops) under its fingerprint; keeps the newest 300."""
    RESULTS.mkdir(parents=True, exist_ok=True)
    dst = RESULTS / key
    tmp = RESULTS / f".{key}.{uuid.uuid4().hex[:6]}"
    try:
        shutil.copytree(art, tmp)
        write_json(tmp / "result.json", dict(schema=RESULT_SCHEMA, rep=rep, photo_meta=photo_meta,
                                             preview_scale=preview_scale, events=events, created=now()))
        if dst.exists():
            shutil.rmtree(dst, ignore_errors=True)
        os.replace(tmp, dst)
    except OSError:
        shutil.rmtree(tmp, ignore_errors=True)
        return
    entries = sorted((p for p in RESULTS.iterdir() if p.is_dir() and not p.name.startswith(".")),
                     key=lambda p: p.stat().st_mtime)
    for old in entries[:-300]:
        shutil.rmtree(old, ignore_errors=True)


def has_result(photo_sha, ga=None, wires=None, checklist=None, enabled=None, engine="auto", prev=None):
    """True when an inspection with these fingerprints can be answered from stored metadata, without the photo."""
    kind = engine_for(ga, engine, prev)
    inputs = {}
    for key, given, default in (("ga", ga, config.DEMO_GA), ("wires", wires, config.DEMO_WIRES),
                                ("checklist", checklist, config.DEMO_CHECKLIST)):
        if given:
            inputs[key] = dict(sha256=sha256_file(given))
        elif kind == "legacy":
            inputs[key] = dict(demo=True, name=default.name)
        else:
            inputs[key] = dict(missing=True)
    return load_result(result_key(photo_sha, inputs, _enabled_for(kind, enabled), kind)) is not None


# ------------------------------------------------------------------ the pipeline
def run_inspection(photo, ga=None, wires=None, checklist=None, enabled=None, parent=None, label=None,
                   source="upload", photo_name=None, input_names=None, forge=None, on_event=None,
                   parent_data=None, reports=True, engine="auto", photo_sha=None, showcase=None):
    """Inspect one photo end to end and store the run (metadata only). Returns the run's meta dict.

    ``parent`` re-checks a stored run; ``parent_data=(meta, report)`` re-checks a run held by the caller
    (the stateless cloud API, where the browser keeps the history). ``reports=False`` skips the Excel and
    PDF so a caller can return results first and call ``build_reports`` afterwards. ``engine`` = auto | legacy | real.
    ``photo=None`` with ``photo_sha`` answers from the result cache or raises NeedPhoto. ``showcase`` (a case of
    showcase.py: title, story, what was planted) is kept with the run for the History page.
    """
    emit = on_event or (lambda ev: None)
    rid = new_id(); d = config.RUNS / rid; d.mkdir(parents=True)
    try:
        return _run(rid, d, photo, ga, wires, checklist, enabled, parent, label, source, photo_name,
                    input_names, forge, emit, parent_data, reports, engine, photo_sha, showcase)
    except Exception:
        shutil.rmtree(d, ignore_errors=True)
        raise


def _engine(kind, photo, inputs, enabled, emit, work, d, rec):
    """Run one engine on the photo; the full-resolution pixels exist only inside this call (and ``work``).

    Returns (report, artefacts folder with preview / issue crops / rectified, preview scale)."""
    art = work / "art"
    (art / "crops").mkdir(parents=True, exist_ok=True)
    if kind == "legacy":
        photo_path = prepare_photo(photo, work)
        rep, ctx = inspect(str(photo_path), inputs["ga"]["path"], inputs["wires"]["path"], enabled=enabled, on_event=emit)
        if inputs["ga"].get("demo") and rep["registration"]["anchors"] == 0:
            raise WrongDrawing(DEMO_MISMATCH.format(n=rep["registration"]["targets"], cab=rep["ga"]["cabinet"]))
        rgb = ctx.photo
        rim = Image.fromarray(ctx.rect)
        k = RECT_PX_PER_MM / R
        rim.resize((max(1, round(rim.width * k)), max(1, round(rim.height * k))), Image.LANCZOS).save(
            art / FILES["rectified"], quality=85, optimize=True)
    else:
        import pipeline
        [(rep, ctx)] = pipeline.inspect_real([photo], inputs["ga"]["path"], wires=inputs["wires"].get("path"),
                                             checklist=inputs["checklist"].get("path"), enabled=enabled, on_event=emit, run_dir=d,
                                             names=[rec["name"]])
        if rep["registration"]["anchors"] == 0 and all(r["result"] == "NOT_IN_VIEW" for r in rep["results"]):
            raise WrongDrawing(REAL_MISMATCH.format(name=inputs["ga"]["name"], cab=rep["ga"]["cabinet"]))
        rgb = ctx.photo["rgb"]
    _, scale = save_preview(rgb, art / FILES["preview"])
    if kind == "real" and rep["registration"].get("localised") and rep["registration"].get("fits"):
        try:                                   # the drawing through the photo's perspective, and the photo straightened
            prev = np.asarray(Image.open(art / FILES["preview"]).convert("RGB"))
            rep["rectify"] = rp.real_overlays(inputs["ga"]["path"], ctx.kb.get("views") or {}, rep["registration"]["fits"],
                                              prev, rep["photo_size"], art)
        except Exception:  # noqa: BLE001 - overlays are a view aid; the inspection result stands without them
            rep["rectify"] = None
    full = Image.fromarray(np.ascontiguousarray(rgb))
    for r in rep["results"]:
        if r["result"] in ("FAIL", "REVIEW") or (kind == "real" and r.get("candidate") in ("CANDIDATE_FAULT", "UNCERTAIN")):
            rp.found_crop(full, rep["H"], r["box"], rp.COL.get(r["result"], rp.COL["FAIL"]), max_side=CROP_SIDE).save(
                art / "crops" / f"{_safe(r['id'])}_found.jpg", quality=86)
    return rep, art, scale


def _install(src, d):
    """Copy a result's stored artefacts (preview, rectified, overlays, crops) into a run folder."""
    for name in (FILES["preview"], FILES["rectified"], FILES["overlay"], FILES["drawing"]):
        if (Path(src) / name).exists():
            shutil.copyfile(Path(src) / name, d / name)
    if (Path(src) / "crops").is_dir():
        shutil.copytree(Path(src) / "crops", d / "crops", dirs_exist_ok=True)


def _run(rid, d, photo, ga, wires, checklist, enabled, parent, label, source, photo_name, input_names, forge, emit,
         parent_data=None, reports=True, engine="auto", photo_sha=None, showcase=None):
    prev_meta = prev_rep = None
    if parent_data:
        prev_meta, prev_rep = parent_data
        parent = prev_meta.get("id") or parent
    elif parent:
        prev_meta = get_meta(parent)
        prev_rep = get_report(parent)
    ga, wires, checklist, input_names = inherited_inputs(prev_meta, ga, wires, checklist, input_names)
    validate_request(ga, engine, source, prev_meta, photo, enabled)
    for key, path in (("ga", ga), ("wires", wires), ("checklist", checklist)):
        if path:
            validate_input(path, key)
    kind = engine_for(ga, engine, (prev_meta or {}).get("engine_kind"))
    enabled = _enabled_for(kind, enabled)
    inputs = resolve_inputs(d, ga, wires, checklist, input_names, defaults=(kind == "legacy"))
    if kind == "real" and not inputs["ga"].get("path"):
        raise InspectionError("The real-cabinet engine needs the cabinet's drawing set (vector PDF with the GA views).")

    work = WORK / rid
    events = []

    def rec_emit(ev):
        events.append(ev)
        emit(ev)
    try:
        if photo is None:
            if not photo_sha:
                raise InspectionError("No photo given.")
            hit = load_result(result_key(photo_sha, inputs, enabled, kind))
            if hit is None:
                raise NeedPhoto("photo")
            rec = hit["photo_meta"]
        else:
            rec = photo_record(photo, photo_name)
            if photo_sha and photo_sha != rec["sha256"]:
                raise InspectionError("The photo does not match its fingerprint (the upload was corrupted). Try again.")
            key = None if forge else result_key(rec["sha256"], inputs, enabled, kind)
            hit = load_result(key)
        if hit:
            rep, scale = hit["rep"], hit["preview_scale"]
            for ev in hit.get("events") or []:
                if ev.get("status") in ("done", "retry"):
                    emit(dict(ev, t=0.0, cached=True))
            _install(hit["dir"], d)
        else:
            work.mkdir(parents=True, exist_ok=True)
            rep, art, scale = _engine(kind, photo, inputs, enabled, rec_emit, work, d, rec)
            if key:
                store_result(key, rep, art, rec, scale, events)
            _install(art, d)
    finally:
        shutil.rmtree(work, ignore_errors=True)

    rep = _guard_registration_report(json.loads(json.dumps(rep)))  # never mutate a cached result in place
    rep["safety_schema"] = SAFETY_SCHEMA
    rep["photo"] = FILES["preview"]
    rep["preview_scale"] = scale
    rep["photo_meta"] = dict(rep.get("photo_meta") or {}, **rec)
    t = datetime.datetime.now(); emit(dict(stage="report", status="start", t=rep["timing"]["total"]))
    previous_ga = ((prev_meta or {}).get("inputs") or {}).get("ga") or {}
    current_ga = inputs["ga"]
    same_reference = bool(prev_rep and kind == prev_meta.get("engine_kind", "legacy") and
                          not _registration_is_invalid(prev_rep) and
                          prev_rep.get("ga", {}).get("cabinet") == rep.get("ga", {}).get("cabinet") and
                          ((previous_ga.get("sha256") and previous_ga.get("sha256") == current_ga.get("sha256")) or
                           (previous_ga.get("demo") and current_ga.get("demo"))))
    if same_reference:
        rp.recheck(prev_rep, rep)
        status_key = "recheck"
    else:
        rp.numbered(rep["results"])
        status_key = "result"
        if prev_rep:
            if _registration_is_invalid(prev_rep):
                code, message = "previous_unverified", "The previous inspection could not be aligned to its drawing. "
            else:
                code, message = "drawing_changed", "The reference drawing or cabinet changed. "
            rep.setdefault("warnings", []).append(dict(code=code, message=message +
                "This is a fresh inspection; previous issues were not automatically closed."))
    rp.markup(str(d / FILES["preview"]), rep, str(d / FILES["marked"]), status_key=status_key, scale=scale)
    if kind == "legacy":
        rp.ga_overlay(inputs["ga"]["path"], rep["ga"]["size"], str(d / FILES["overlay"]))
    calibration.apply(rep)                    # calibrated confidence per item and for the photo
    c = rep["counts"]
    judged = len(rep["results"]) - c.get("NOT_IN_VIEW", 0)
    created = now()
    version = (prev_meta.get("version", 1) + 1) if prev_meta else 1
    meta = dict(id=rid, created=created, label=label or (f"Re-check v{version}" if prev_meta else "Inspection"),
                source=source, photo_name=photo_name or rec.get("name") or "photo", photo=FILES["preview"],
                photo_meta=rec, storage="metadata", inputs=inputs, enabled=enabled, cabinet=rep["ga"]["cabinet"],
                counts=c, items=len(rep["results"]),
                verdict=verdict(c, judged, rep.get("warnings"), rep), timing=rep["timing"],
                warnings=rep.get("warnings", []),
                registration={k: rep["registration"].get(k) for k in
                              ("anchors", "inliers", "outliers", "refined", "targets", "valid", "invalid_reason")},
                issues=issue_rows(rep), confidence=rep.get("confidence"),
                parent=parent, root=(prev_meta.get("root") or parent) if prev_meta else rid, version=version,
                recheck=rep.get("recheck"), engine=config.VERSION, engine_kind=kind, cached=bool(hit))
    if kind == "real":
        meta.update(coverage=rep.get("coverage"), views=rep.get("views"),
                    agent={k: (rep.get("agent") or {}).get(k) for k in ("enabled", "backend", "model", "calls",
                                                                       "cost_usd", "reason")})
    if forge:
        forge["score"] = score_forge(rep, forge["faults"])
        meta["forge"] = forge
    if showcase:
        meta["showcase"] = showcase
    write_json(d / FILES["json"], rep)
    if reports:
        _reports(rid, rep, meta)
    meta["timing"] = dict(rep["timing"], report=round((datetime.datetime.now() - t).total_seconds(), 3))
    write_json(d / "meta.json", meta)
    emit(dict(stage="report", status="done", t=rep["timing"]["total"] + meta["timing"]["report"],
              files=[f for f in FILES.values() if (d / f).exists()], storage="metadata"))
    if rep.get("recheck"):
        emit(dict(stage="recheck", status="done", t=rep["timing"]["total"] + meta["timing"]["report"], **rep["recheck"]))
    return meta


def _checklist_ids(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        try:
            return {str(r[0]) for r in wb.active.iter_rows(min_row=2, values_only=True) if r and r[0]}
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        return set()


def _reports(rid, rep, meta):
    d = run_dir(rid)
    template = input_path(rid, meta, "checklist")
    generated = d / "checklist_template.xlsx"
    if meta.get("engine_kind") == "real" and (template is None or
                                               not _checklist_ids(template) & {r["id"] for r in rep["results"]}):
        template = rp.checklist_template(rep, str(generated))
    try:
        rp.fill_checklist(str(template), rep, str(d / FILES["checklist"]), meta=meta)
    finally:
        generated.unlink(missing_ok=True)
    pdf_report.build_pdf(str(d / FILES["pdf"]), rep, meta, str(d / FILES["marked"]),
                         lambda r: (item_crop(rid, r["id"], "expected", rep=rep, meta=meta),
                                    item_crop(rid, r["id"], "found", rep=rep, meta=meta)))


def build_reports(rid):
    """PDF report + filled Excel checklist for a run stored with ``reports=False``."""
    meta, rep = get_meta(rid), get_report(rid)
    if meta.get("historical_guard"):
        raise InspectionError("This archived inspection contains unverified conclusions. Start a new inspection with "
                              "the matching drawing to generate current PDF and checklist exports.")
    _reports(rid, rep, meta)
    return artefact(rid, "pdf"), artefact(rid, "checklist")


def save_review(rid, item_id, decision, note="", by="Inspector"):
    """Record a human decision on an item (REVIEW items mostly): the audit trail keeps who, what and when."""
    if decision not in ("PASS", "FAIL", "CLEAR"):
        raise ValueError("decision must be PASS, FAIL or CLEAR")
    d = run_dir(rid)
    meta = get_meta(rid)
    rep = get_report(rid)
    if not any(r["id"] == item_id for r in rep["results"]):
        raise KeyError(item_id)
    reviews = meta.setdefault("reviews", {})
    if decision == "CLEAR":
        reviews.pop(item_id, None)
    else:
        reviews[item_id] = dict(decision=decision, note=str(note)[:500], by=str(by)[:80], at=now())
    write_json(d / "meta.json", meta)
    return meta


# ------------------------------------------------------------------ reading runs
def _registration_is_invalid(rep):
    registration = rep.get("registration") or {}
    return bool(registration.get("valid") is False or registration.get("localised") is False or any(
        w.get("code") == "weak_registration" for w in rep.get("warnings") or []))


def _guard_registration_report(rep, meta=None):
    """Apply current evidence boundaries to stored reports without altering the original audit files."""
    kind = rep.get("engine_kind") or (meta or {}).get("engine_kind", "legacy")
    if not (_registration_is_invalid(rep) or (meta and _registration_is_invalid(meta))):
        return _guard_real_item_report(rep, meta) if kind == "real" else rep
    rep["engine_kind"] = kind
    warnings = rep.setdefault("warnings", [])
    weak = next((w for w in warnings if w.get("code") == "weak_registration"), None)
    if weak is None:
        weak = {"code": "weak_registration", "message": "The photo was not reliably placed on its reference drawing. "
                "Confirm the matching drawing and capture a sharp, straight-on cabinet view."}
        warnings.append(weak)
    if kind == "legacy":
        rep["results"] = unregistered_results(rep["ga"], rep.get("wires") or [], rep.get("enabled") or ORDER,
                                               weak["message"])
    else:
        rows = []
        for item in rep.get("results") or []:
            if item.get("id") == "REG-DRAWING":
                continue
            row = dict(item)
            for key in ("marker", "recheck", "prev_result", "prev_found"):
                row.pop(key, None)
            if row.get("result") != "NOT_IN_VIEW":
                row.update(result="SKIP", found="Unverified: photo not aligned with drawing", conf=0.0,
                           verified_by=[], guarded=True, root_cause="REG-DRAWING")
                evidence = list(row.get("evidence") or [])
                row["evidence"] = evidence if weak["message"] in evidence else evidence + [weak["message"]]
            rows.append(row)
        size = rep.get("photo_size") or rep.get("ga", {}).get("size") or [1, 1]
        rows.insert(0, dict(id="REG-DRAWING", check="REGISTRATION", target="Photo / drawing alignment",
                            expected="Matching cabinet and usable alignment", found="Confirm matching drawing and capture a clear cabinet view",
                            result="REVIEW", box=[0, 0, *size], conf=0.0, evidence=[weak["message"]],
                            root_cause=True, verified_by=[], view=None, drawing=None))
        rep["results"] = rows
        cov = rep.get("coverage") or {}
        cov.update(verified=0, ratio=0.0, complete=False, approval_allowed=False,
                   unresolved=len(rows), in_view=0, not_in_view=cov.get("total", 0))
        for view in (cov.get("by_view") or {}).values():
            view["verified"] = 0
        rep["coverage"] = cov
    rep["counts"] = counts(rep["results"])
    if kind == "real":
        rep["counts"]["NOT_IN_VIEW"] = sum(r["result"] == "NOT_IN_VIEW" for r in rep["results"])
    rep.setdefault("registration", {}).update(valid=False, invalid_reason=weak["message"])
    if kind == "real":
        rep["registration"]["localised"] = False
    rep.pop("recheck", None)
    rep["registration_guard"] = True
    rp.numbered(rep["results"])
    return rep


def _guard_real_item_report(rep, meta=None):
    """Old global fits and model answers do not verify a component in an unfitted drawing view."""
    import consensus
    registration = rep.get("registration") or (meta or {}).get("registration") or {}
    changed = []
    for row in rep.get("results") or []:
        reason = consensus.conclusive_evidence_issue(row, registration)
        if reason is None:
            continue
        row.update(result="REVIEW", found=f"Unverified: {row.get('found', '')}", verified_by=[], conf=0.0,
                   guarded=True, hint="Review the local evidence and capture this drawing view with readable surrounding labels")
        evidence = list(row.get("evidence") or [])
        row["evidence"] = evidence if reason in evidence else evidence + [reason]
        changed.append(row["id"])
    coverage = rep.get("coverage")
    missing_coverage = not isinstance(coverage, dict) or not all(k in coverage for k in (
        "total", "verified", "complete", "approval_allowed"))
    if not changed and not missing_coverage:
        return rep
    rep["engine_kind"] = "real"
    rows = rep.get("results") or []
    for row in rows:
        for key in ("marker", "recheck", "prev_result", "prev_found"):
            row.pop(key, None)
    verified = {row["target"] for row in rows if row.get("check") == "LABEL" and row["result"] in ("PASS", "FAIL")}
    cov = dict(coverage or {})
    total = cov.get("total") or (rep.get("kb") or {}).get("tags") or 0
    total = total if isinstance(total, (int, float)) and total >= 0 else 0
    cov.update(verified=len(verified), total=total, ratio=round(len(verified) / total, 3) if total else 0.0,
               complete=False, approval_allowed=False, unresolved=sum(r["result"] not in ("PASS", "FAIL") for r in rows))
    for view, detail in (cov.get("by_view") or {}).items():
        detail["verified"] = len({r["target"] for r in rows if r.get("check") == "LABEL" and r.get("view") == view
                                  and r["result"] in ("PASS", "FAIL")})
    rep["coverage"] = cov
    rep["counts"] = consensus.counts(rows)
    rep.pop("recheck", None)
    warning = {"code": "historical_evidence", "message": "Stored conclusions were reassessed using local OCR and "
               "per-view registration. Missing coverage or unverified model judgments cannot approve the cabinet."}
    if not any(w.get("code") == warning["code"] for w in rep.get("warnings") or []):
        rep.setdefault("warnings", []).append(warning)
    rep["evidence_guard"] = True
    rp.numbered(rows)
    return rep


def issue_rows(rep):
    """The numbered issues of a report as stored in the run's metadata, with where each belongs on the assembly line
    (device family, cabinet area, line station) and its calibrated confidence, for QC analytics."""
    rows = []
    for r in rep["results"]:
        if not r.get("marker"):
            continue
        dims = {k: r[k] for k in ("family", "area", "station") if r.get(k)} or assembly.dims(r, rep)
        rows.append(dict(id=r["id"], result=r["result"], found=r["found"], marker=r.get("marker"),
                         recheck=r.get("recheck"), check=r.get("check"), target=r.get("target"),
                         family=dims.get("family"), area=dims.get("area"), station=dims.get("station"),
                         confidence=r.get("confidence")))
    return rows


def get_meta(rid):
    meta = read_json(run_dir(rid) / "meta.json")
    rep = get_report(rid, meta=meta)
    if rep.get("registration_guard") or rep.get("evidence_guard"):
        meta.update(counts=rep["counts"], items=len(rep["results"]),
                    verdict=verdict(rep["counts"], len(rep["results"]), rep.get("warnings"), rep),
                    issues=issue_rows(rep), warnings=rep["warnings"],
                    historical_guard=rep.get("safety_schema") != SAFETY_SCHEMA, recheck=None)
        if rep.get("registration_guard"):
            meta.setdefault("registration", {}).update(valid=False, invalid_reason=rep["registration"]["invalid_reason"])
        if rep.get("engine_kind") == "real":
            if rep.get("registration_guard"):
                meta["registration"]["localised"] = False
            meta["coverage"] = rep["coverage"]
    return meta


def get_report(rid, meta=None):
    rep = read_json(run_dir(rid) / FILES["json"])
    if meta is None:
        metadata_path = run_dir(rid) / "meta.json"
        meta = read_json(metadata_path) if metadata_path.exists() else None
    return _guard_registration_report(rep, meta)


def list_runs(limit=200):
    metas = []
    for p in config.RUNS.glob("*/meta.json"):
        try:
            metas.append(get_meta(p.parent.name))
        except (OSError, ValueError, KeyError, TypeError):
            continue
    metas.sort(key=lambda m: m["id"], reverse=True)
    return metas[:limit]


def chain(rid):
    """The run's re-check chain, oldest first (v1 -> v2 -> ...), including later re-checks of it."""
    meta = get_meta(rid)
    root = meta.get("root") or rid
    runs = [m for m in list_runs(10_000) if (m.get("root") or m["id"]) == root]
    return sorted(runs, key=lambda m: (m.get("version", 1), m["id"]))


def run_payload(rid):
    """Everything the UI needs to draw a run: meta, results with photo-space quads, drawing layout."""
    meta, rep = get_meta(rid), get_report(rid)
    H = rep["H"]
    q = lambda box: np.round(rp.to_photo(box, H), 1).tolist()
    for r in rep["results"]:
        r["quad"] = q(r["box"])
    layout = [dict(tag=d["tag"], kind=d["type"], box=d["box"], quad=q(d["box"]), label_box=d["label_box"])
              for d in rep["ga"]["devices"]]
    layout += [dict(tag=t["tag"], kind="terminal", box=t["box"], quad=q(t["box"])) for t in rep["ga"]["terminals"]]
    Wmm, Hmm = rep["ga"]["size"]
    out = dict(meta=meta, results=rep["results"], layout=layout, registration=rep["registration"],
               photo_size=rep.get("photo_size"), size_mm=[Wmm, Hmm], backplate_quad=q((0, 0, Wmm, Hmm)),
               rails=rep["ga"]["rails"], wires=rep.get("wires", []), timing=rep.get("timing"),
               files={k: v for k, v in FILES.items()}, rectified=rep.get("rectified", True),
               engine_kind=rep.get("engine_kind", "legacy"))
    if rep.get("engine_kind") == "real":
        out.update(coverage=rep.get("coverage"), views=rep.get("views"), agent=rep.get("agent"), ocr=rep.get("ocr"),
                   quality=rep.get("quality"), notes=rep.get("notes") or [], photo_meta=rep.get("photo_meta"))
    out["wire_rule"] = rep.get("wire_rule")
    out["rectify"] = rep.get("rectify")
    return out


def artefact(rid, name):
    d = run_dir(rid)
    if name in ("photo", "photo.jpg", "photo.png", "preview"):
        return d / get_meta(rid)["photo"]
    fname = FILES.get(name, name)
    if fname not in FILES.values():
        raise KeyError(name)
    if fname == FILES["json"]:
        # Derived download: keep report.json as the immutable original audit record.
        safe = d / "report.safe.json"
        write_json(safe, get_report(rid))
        return safe
    if fname in {FILES[k] for k in ("pdf", "checklist", "marked", "rectified", "overlay")} and get_meta(rid).get("historical_guard"):
        raise InspectionError("This archived export contains unverified inspection conclusions and is unavailable. "
                              "Start a new inspection with the matching drawing to generate safe exports. "
                              "The guarded JSON report and original photo preview remain available.")
    p = d / fname
    if not p.exists():
        raise KeyError(name)
    return p


def _text_card(lines, size=(520, 300)):
    """A small proof image with text (items that have no drawing position, e.g. a label not in the drawing set)."""
    from PIL import ImageDraw, ImageFont
    im = Image.new("RGB", size, (255, 255, 255))
    dr = ImageDraw.Draw(im)
    f = ImageFont.truetype(config.font("DejaVuSans-Bold.ttf"), 20)
    fs = ImageFont.truetype(config.font("DejaVuSans.ttf"), 15)
    y = 40
    for i, ln in enumerate(lines):
        dr.text((24, y), ln, fill=(31, 79, 168) if i == 0 else (60, 70, 85), font=f if i == 0 else fs)
        y += 34 if i == 0 else 24
    return im


def _drawing_crop(rid, meta, r):
    """Expected (GA drawing) crop of a real-engine item, made from the stored drawing set."""
    from kb import drawing_kb
    pdf = input_path(rid, meta, "ga")
    dr = r.get("drawing") or {}
    if not pdf or not dr.get("view"):
        if r["check"] == "UNKNOWN_LABEL":
            return _text_card([f"'{r['target']}' is not in the drawing", "No device with this tag exists on any GA view,",
                               "in the nameplate list or in the schematics."])
        if r["check"] == "FERRULE":
            return _text_card([f"Ferrule {r['target']}", "Checked against the wiring extracted from",
                               "the drawing set's schematics."])
        return _text_card([str(r["target"]), "No drawing position for this item."])
    kb = drawing_kb.build_kb(str(pdf))
    tag = r["target"] if r["check"] in ("MISSING", "LABEL") else None
    try:
        im = drawing_kb.drawing_crop(kb, str(pdf), dr["view"], tag=tag, bbox=None if tag else dr.get("bbox"),
                                     margin_pt=24, dpi=600)
    except (KeyError, ValueError, TypeError):
        return _text_card([str(r["target"]), "Not drawn on this view."])
    im.thumbnail((int(CROP_SIDE * 1.4), int(CROP_SIDE * 1.4)), Image.LANCZOS)
    return im


def item_crop(rid, item_id, kind, rep=None, meta=None):
    """Proof image for one item: 'expected' (GA crop), 'found' (photo crop) or 'aligned' (rectified crop).

    Photo crops of issues were cut from the full-resolution photo during the inspection; others come from the stored
    preview (the full photo is not kept)."""
    d = run_dir(rid)
    if kind not in ("expected", "found", "aligned"):
        raise KeyError(kind)
    out = d / "crops" / f"{_safe(item_id)}_{kind}.{'png' if kind == 'expected' else 'jpg'}"
    if out.exists():
        return out
    rep = rep or get_report(rid); meta = meta or get_meta(rid)
    r = next((x for x in rep["results"] if x["id"] == item_id), None)
    if r is None:
        raise KeyError(item_id)
    real = meta.get("engine_kind") == "real"
    out.parent.mkdir(exist_ok=True)
    color = rp.COL.get(r.get("recheck") or r["result"], rp.COL["FAIL"])
    if kind == "expected":
        if real:
            _drawing_crop(rid, meta, r).save(out)
        else:
            png, s = ga_raster(input_path(rid, meta, "ga"))
            rp.expected_crop(Image.open(png), s, r["box"]).save(out)
    elif kind == "found":
        src = d / meta["photo"]
        scale = float(rep.get("preview_scale") or 1.0) if meta["photo"] == FILES["preview"] else 1.0
        rp.found_crop(Image.open(src), rep["H"], r["box"], color, scale=scale, max_side=CROP_SIDE).save(out, quality=88)
    else:
        if real or not (d / FILES["rectified"]).exists():
            raise KeyError(kind)
        rect = Image.open(d / FILES["rectified"])
        k = rect.width / float(rep["ga"]["size"][0])
        rp.aligned_crop(rect, r["box"], color, k=k).save(out, quality=88)
    return out


def delete_runs():
    """Delete every stored run and the cached results (fingerprint cache): nothing of the photos stays behind."""
    n = 0
    for p in config.RUNS.iterdir():
        if p.is_dir() and _ID.match(p.name):
            shutil.rmtree(p, ignore_errors=True); n += 1
    shutil.rmtree(RESULTS, ignore_errors=True)
    return n


# ------------------------------------------------------------------ samples
def samples():
    mf = config.SAMPLES / "manifest.json"
    items = read_json(mf) if mf.exists() else []
    return [dict(s, id=s["file"].split("_")[0], fault_text=[describe_fault(f) for f in s.get("faults", [])])
            for s in items if (config.SAMPLES / s["file"]).exists()]


def sample_path(name):
    for s in samples():
        if name in (s["file"], s["id"]):
            return config.SAMPLES / s["file"]
    raise KeyError(name)


# ------------------------------------------------------------------ Fault Forge
def forge_photo(dst, k=None, strength=1.0, seed=None, faults=None):
    """Render a never-seen cabinet with hidden planted faults. Returns (path, faults, seed)."""
    seed = int(seed) if seed is not None else random.randrange(10_000, 99_999)
    rng = np.random.default_rng(seed)
    faults = faults if faults is not None else random_faults(rng, None if k is None else int(k))
    photo, _ = render(faults, seed=seed, strength=float(strength))
    photo.save(dst, quality=90)
    return Path(dst), faults, seed


def score_forge(rep, faults):
    res = {r["id"]: r for r in rep["results"]}
    planted = []
    for f in faults:
        item = expected_item(f); r = res.get(item)
        planted.append(dict(fault=f, text=describe_fault(f), item=item, caught=bool(r and r["result"] == "FAIL"),
                            result=r["result"] if r else None, found=r["found"] if r else None))
    exp = {p["item"] for p in planted}
    fa = [r["id"] for r in rep["results"] if r["result"] == "FAIL" and r["id"] not in exp]
    rv = [r["id"] for r in rep["results"] if r["result"] == "REVIEW" and r["id"] not in exp]
    skip = sum(r["result"] == "SKIP" for r in rep["results"])
    return dict(planted=planted, caught=sum(p["caught"] for p in planted), n_planted=len(planted),
                false_alarms=fa, reviews=rv, correct_items=len(rep["results"]) - skip - len(planted))


def run_forge(k=None, strength=1.0, seed=None, faults=None, on_event=None, **kw):
    tmp = CACHE / f"forge_{uuid.uuid4().hex[:8]}.jpg"
    try:
        path, faults, seed = forge_photo(tmp, k=k, strength=strength, seed=seed, faults=faults)
        if on_event:
            on_event(dict(stage="forge", status="done", seed=seed, planted=len(faults)))
        return run_inspection(path, label=f"Fault Forge #{seed}", source="forge", photo_name=f"forge_{seed}.jpg",
                              forge=dict(seed=seed, strength=round(float(strength), 2), faults=faults), on_event=on_event,
                              **kw)
    finally:
        tmp.unlink(missing_ok=True)


def forge_live():
    """Running score over every Fault Forge cabinet inspected on this machine."""
    runs = [m for m in list_runs(10_000) if m.get("forge", {}).get("score")]
    sc = [m["forge"]["score"] for m in runs]
    return dict(cabinets=len(runs), planted=sum(s["n_planted"] for s in sc), caught=sum(s["caught"] for s in sc),
                false_alarms=sum(len(s["false_alarms"]) for s in sc), reviews=sum(len(s["reviews"]) for s in sc),
                correct_items=sum(s["correct_items"] for s in sc),
                runs=[dict(id=m["id"], seed=m["forge"]["seed"], created=m["created"], planted=m["forge"]["score"]["n_planted"],
                           caught=m["forge"]["score"]["caught"], false_alarms=len(m["forge"]["score"]["false_alarms"]),
                           reviews=len(m["forge"]["score"]["reviews"])) for m in runs[:40]])


def selftest():
    """Saved results of the offline Fault Forge self-test (python eval_forge.py 50 200)."""
    d = config.OUT / "forge"
    summary = read_json(d / "summary.json") if (d / "summary.json").exists() else None
    rows = []
    if (d / "results.json").exists():
        for r in read_json(d / "results.json"):
            rows.append(dict(i=r["i"], strength=round(r["strength"], 2), planted=len(r["faults"]),
                             caught=sum(r["caught"].values()), false_alarms=len(r["false_alarms"]),
                             reviews=len(r["reviews"]), anchors=r.get("anchors"), sec=r.get("sec"),
                             faults=[describe_fault(f) for f in r["faults"]]))
        rows.sort(key=lambda r: r["i"])
    return dict(summary=summary, cabinets=rows)


def checks_info():
    return [dict(name=n, title=getattr(c, "title", n), question=getattr(c, "question", ""),
                 doc=(c.__doc__ or "").strip(), builtin=n in ORDER) for n, c in CHECKS.items()]
