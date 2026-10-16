"""Run store, re-check chain, artefacts and Fault Forge scoring."""
import zipfile

from openpyxl import load_workbook

import config
import service


def sample(sid):
    return next(config.SAMPLES.glob(f"{sid}_*.jpg"))


def test_inspect_then_recheck_closes_three_and_keeps_one_open(runs_dir):
    v1 = service.run_inspection(sample("01"), source="sample")
    assert v1["verdict"]["label"] == "HOLD" and v1["counts"]["FAIL"] == 4
    assert [i["marker"] for i in v1["issues"]] == [1, 2, 3, 4]

    v2 = service.run_inspection(sample("05"), parent=v1["id"], source="sample")
    assert v2["version"] == 2 and v2["root"] == v1["id"]
    assert v2["recheck"] == {"CLOSED": 3, "OPEN": 1, "REGRESSION": 0, "NEW": 0}
    still_open = [i for i in v2["issues"] if i["recheck"] == "OPEN"]
    assert [(i["id"], i["marker"]) for i in still_open] == [("FER-X3:3", 4)], "an open item keeps its number"
    assert [m["id"] for m in service.chain(v2["id"])] == [v1["id"], v2["id"]]


def test_recheck_flags_regressions(runs_dir):
    v1 = service.run_inspection(sample("02"))                       # clean cabinet, all PASS
    v2 = service.run_inspection(sample("04"), parent=v1["id"])      # PS11 mislabelled, CB42 removed
    assert {i["id"]: i["recheck"] for i in v2["issues"]} == {"GA-CB42": "REGRESSION", "LBL-PS11": "REGRESSION"}


def test_every_artefact_is_written(runs_dir):
    m = service.run_inspection(sample("03"))
    for name in ("marked", "rectified", "overlay", "json", "checklist", "pdf", "photo"):
        assert service.artefact(m["id"], name).stat().st_size > 0
    assert service.artefact(m["id"], "pdf").read_bytes()[:4] == b"%PDF"
    assert zipfile.is_zipfile(service.artefact(m["id"], "checklist"))
    wb = load_workbook(service.artefact(m["id"], "checklist"))
    assert wb.sheetnames[0] == "Summary"
    rows = {r[0]: r for r in wb["QC Checklist"].iter_rows(min_row=2, values_only=True)}
    assert rows["FER-X3:5"][5] == "FAIL" and rows["COL-X3:7"][4] == "BK" and rows["GA-PS11"][5] == "PASS"
    for kind in ("expected", "found", "aligned"):
        assert service.item_crop(m["id"], "FER-X3:5", kind).exists()


def test_run_payload_projects_every_item_onto_the_photo(runs_dir):
    m = service.run_inspection(sample("01"))
    p = service.run_payload(m["id"])
    W, H = p["photo_size"]
    assert len(p["layout"]) == 20 and len(p["results"]) == 41
    for r in p["results"]:
        assert len(r["quad"]) == 4 and all(0 <= x <= W and 0 <= y <= H for x, y in r["quad"])
    outliers = {a["tag"] for a in p["registration"]["anchor_points"] if a["inlier"] is False}
    assert outliers == {"CB43", "CB44"}, "the swapped breakers' labels do not fit the drawing"


def test_fault_forge_scores_against_the_hidden_answer_key(runs_dir):
    faults = [dict(type="missing_component", tag="PS11"), dict(type="wrong_colour", n=2, colour="RD")]
    m = service.run_forge(seed=4242, strength=1.0, faults=faults)
    s = m["forge"]["score"]
    assert s["n_planted"] == 2 and s["caught"] == 2 and s["false_alarms"] == []
    live = service.forge_live()
    assert live["cabinets"] == 1 and live["caught"] == 2


def test_delete_runs(runs_dir):
    service.run_inspection(sample("02"))
    assert service.delete_runs() == 1 and service.list_runs() == []


def test_recheck_reuses_an_uploaded_drawing_set(runs_dir, tmp_path):
    import shutil
    ga = tmp_path / "customer_GA.pdf"
    shutil.copyfile(config.DEMO_GA, ga)
    v1 = service.run_inspection(sample("01"), ga=ga, input_names={"ga": "customer_GA.pdf"})
    assert v1["inputs"]["ga"]["name"] == "customer_GA.pdf" and not v1["inputs"]["ga"]["demo"]
    v2 = service.run_inspection(sample("05"), parent=v1["id"])        # no inputs given: inherit v1's
    assert v2["inputs"]["ga"]["name"] == "customer_GA.pdf" and not v2["inputs"]["ga"]["demo"]
    assert v2["recheck"]["CLOSED"] == 3
    assert service.item_crop(v2["id"], "FER-X3:3", "expected").exists()
